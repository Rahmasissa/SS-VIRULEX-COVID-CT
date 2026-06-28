from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/samehissa/Downloads/SS-VIRULEX-workspace")
OUT_DOCX = ROOT / "SS_VIRULEX_Revised_Thesis.docx"
OUT_CHANGELOG = ROOT / "SS_VIRULEX_Revised_Thesis_Change_Log.md"


DESCRIPTORS = [
    "mean",
    "std_dev",
    "variance",
    "median",
    "range",
    "skewness",
    "kurtosis",
    "entropy",
    "energy",
    "contrast",
    "mean_abs_dev",
    "min_value",
    "max_value",
    "iqr",
    "percentile_25",
    "percentile_50",
    "percentile_75",
    "signal_to_noise",
    "coef_of_var",
    "autocorrelation",
    "shannon_entropy",
    "root_mean_square",
    "harmonic_mean",
    "geometric_mean",
    "std_error_mean",
    "median_abs_dev",
]


CONCEPTS = [
    "peripheral_ground_glass_opacities",
    "bilateral_involvement",
    "multilobar_distribution",
    "crazy_paving_pattern",
    "absence_of_lobar_consolidation",
    "localized_or_diffuse_presentation",
    "increased_density_in_the_lung",
    "ground_glass_appearance",
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.bold = bold


def add_table(doc: Document, headers: list[str], rows: list[list[str]], caption: str | None = None):
    if caption:
        p = doc.add_paragraph()
        p.style = "Caption"
        p.add_run(caption).bold = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_row = table.rows[0]
    for i, h in enumerate(headers):
        set_cell_text(header_row.cells[i], h, bold=True)
        set_cell_shading(header_row.cells[i], "D9EAF7")
        header_row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
    return table


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = "Caption"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)


def add_figure(doc: Document, image_path: Path, caption: str) -> None:
    if not image_path.exists():
        p = doc.add_paragraph()
        p.style = "Caption"
        p.add_run(f"{caption} [Image file not found: {image_path.name}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(5.6))
    add_caption(doc, caption)


def set_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in [
        ("Title", 20, "1F4E79"),
        ("Heading 1", 16, "1F4E79"),
        ("Heading 2", 13, "1F4E79"),
        ("Heading 3", 11.5, "1F4E79"),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.paragraph_format.space_after = Pt(6)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    set_styles(doc)


def chapter(doc: Document, number: str, title: str) -> None:
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Chapter {number}")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(31, 78, 121)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(31, 78, 121)


def heading(doc: Document, text: str, level: int = 2) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line, size, bold in [
        ("Media Engineering and Technology Faculty", 14, True),
        ("German University in Cairo", 14, True),
        ("", 12, False),
        ("XAI In Clinical Diagnoses", 20, True),
        ("Bachelor Thesis", 16, True),
        ("", 12, False),
        ("Author: Rahma Sameh Issa", 12, False),
        ("Supervisors: Dr. Walaa Gad", 12, False),
        ("Submission Date: Day Month, Year", 12, False),
    ]:
        run = p.add_run(line + "\n")
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.bold = bold
    doc.add_page_break()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Academic Integrity Statement").bold = True
    doc.add_paragraph(
        "This is to certify that: (i) the thesis comprises only my original work toward the Bachelor Degree; "
        "(ii) due acknowledgement has been made in the text to all other material used."
    )
    doc.add_paragraph("Rahma Sameh Issa")
    doc.add_paragraph("Day Month, Year")
    doc.add_page_break()

    doc.add_heading("Acknowledgments", level=1)
    doc.add_paragraph("This section is reserved for the author's acknowledgments.")
    doc.add_page_break()

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "This thesis revises and grounds SS-VIRULEX as an implementation-specific hybrid framework for "
        "explainable COVID-CT image classification. The implemented pipeline combines a SVIS-RULEX-style "
        "statistical descriptor branch with a Med-MICN concept and neural-symbolic branch. The final fused "
        "model feature vector contains 26 statistical descriptors, 8 concept scores, and 2 Med-MICN diagnosis "
        "probabilities, producing 36 model features."
    )
    doc.add_paragraph(
        "The corrected thesis reports the official 746-image COVID-CT split, train-only offline augmentation "
        "for the statistical branch, ZFMIS-style feature filtering and ranking, Decision Tree and RuleFit rule "
        "artifacts, calibrated stronger classifiers, and selected SFMOV/concept-aware SFMOV visual outputs. "
        "Unsupported claims about clinical validation, deployment readiness, external validation, and runtime "
        "GPT-4V/BioViL generation are removed or qualified."
    )
    doc.add_page_break()


def contents(doc: Document) -> None:
    doc.add_heading("Contents", level=1)
    entries = [
        "1 Introduction",
        "1.1 Motivation",
        "1.2 Problem Statement",
        "1.3 Objectives",
        "2 Background",
        "2.1 Concept Overview",
        "2.1.1 Deep Learning and Feature Extraction for Medical Image Analysis",
        "2.1.2 Statistical Feature Engineering",
        "2.1.3 Augmented ZFMIS (Zero-based Filtering with Mutual Importance Selection)",
        "2.1.4 Rule-based Interpretability: Decision Trees and RuleFit",
        "2.1.5 Statistical Feature Map Overlay Visualisation (SFMOV) and Its Concept-Aware Extension",
        "2.1.6 Concept Bottleneck Models and Concept-Based Interpretability",
        "2.1.7 Large Multimodal Models and Vision-Language Models for Medical Concept Semantics",
        "2.1.8 Neural-Symbolic Reasoning and Fuzzy Logic Rules",
        "2.1.9 Late Fusion of Heterogeneous Evidence Sources",
        "2.1.10 Predictive Classifiers, Probability Calibration, and Validation-Based Threshold Selection",
        "2.1.11 A Taxonomy of Explainability Modalities",
        "2.2 Literature Review",
        "2.2.1 Pillar I: Post-Hoc Visual and Attribution Methods",
        "2.2.2 Pillar II: Statistical and Rule-Based Interpretability Frameworks",
        "2.2.3 Pillar III: Concept-based and Neural-Symbolic Ante-Hoc Models",
        "2.2.4 Pillar IV: Hybrid Evidence Fusion and Calibrated Clinical Decision Support",
        "3 Methodology",
        "3.1 Introduction to the Proposed Methodology",
        "3.2 SS-VIRULEX Framework Overview",
        "3.3 Dataset Preparation and Experimental Setting",
        "3.4 Shared Metadata Construction and Preprocessing",
        "3.5 Statistical Branch Based on SVIS-RULEX Principles",
        "3.6 Med-MICN Semantic and Neural-Symbolic Branch",
        "3.7 Statistical-Semantic Feature Fusion",
        "3.8 Augmented ZFMIS Feature Selection",
        "3.9 Rule-Based Classification and Explanation",
        "3.10 Calibrated Stronger Classifiers and Validation-Based Threshold Selection",
        "3.11 Visual Explanation Workflow",
        "3.12 Implementation Details, and Reproducibility",
        "4 Testing and Results",
        "4.1 Dataset and Experimental Setup",
        "4.2 Branch-Level Outputs",
        "4.3 Fused Feature Construction and ZFMIS Results",
        "4.4 Rule-Based Explanation Results",
        "4.5 Visual Explanation Results",
        "4.6 Final Output Reporting",
        "4.7 Final Comparative Results",
        "5 Conclusion",
        "6 Future Work",
        "Appendix",
        "References",
    ]
    for entry in entries:
        doc.add_paragraph(entry)


def chapter1(doc: Document) -> None:
    chapter(doc, "1", "Introduction")
    heading(doc, "1.1 Motivation")
    para(
        doc,
        "The motivation for this thesis is the need to make medical-image classification outputs more "
        "auditable and interpretable while remaining faithful to the actual implementation. The work is "
        "grounded in two reference methods. SVIS-RULEX, as described by Ullah et al. (2025), combines "
        "deep feature extraction, statistical feature engineering, ZFMIS feature selection, Decision Tree "
        "and RuleFit rule extraction, and SFMOV visual overlays. Med-MICN, as described by Hu et al. "
        "(2024), introduces concept prediction, concept embeddings, neural-symbolic reasoning, and "
        "multi-dimensional explanation alignment for medical classification."
    )
    para(
        doc,
        "In the current workspace, SS-VIRULEX adapts these two ideas for COVID-CT classification. The "
        "implementation does not reproduce every reference-paper capability. Instead, it constructs a "
        "practical fused feature pipeline using statistical descriptors, Med-MICN concept scores, and "
        "Med-MICN diagnosis probabilities. The implemented outputs include predictive metrics, concept "
        "metrics, ZFMIS rankings, Decision Tree and RuleFit rule artifacts, calibrated classifier results, "
        "and selected SFMOV/concept-aware SFMOV visual explanations."
    )
    para(
        doc,
        "This revised thesis therefore avoids claims of clinical deployment readiness, clinical validation, "
        "or comparison with unrelated external XAI methods. Its scope is the implemented workspace, "
        "the COVID-CT dataset source, the two core reference methods, and implementation-supported "
        "method/software sources."
    )

    heading(doc, "1.2 Problem Statement")
    para(
        doc,
        "The implemented problem is not to prove a complete clinical decision-support system, but to "
        "investigate whether semantic concept scores and neural-symbolic diagnosis probabilities from "
        "Med-MICN can be fused with statistical descriptors and rule-extraction methods derived from "
        "SVIS-RULEX for explainable COVID-CT classification."
    )
    para(
        doc,
        "The workspace implements late fusion of image-level statistical descriptors, concept scores, "
        "`med_task_prob_covid`, and `med_neural_prob_covid`. It also generates selected concept-aware "
        "SFMOV overlays. Full concept-specific spatial saliency alignment for every concept is not saved "
        "in the current workspace and is treated as a limitation rather than an implemented result."
    )
    para(
        doc,
        "The corrected research question is: How can the implemented Med-MICN concept and neural-symbolic "
        "outputs be combined with SVIS-RULEX-style statistical descriptors, ZFMIS feature selection, rule "
        "extraction, and calibrated classification to produce a grounded explainable COVID-CT classification "
        "pipeline?"
    )

    heading(doc, "1.3 Objectives")
    objectives = [
        "Prepare the official COVID-CT split and train-only augmented statistical-branch manifest.",
        "Extract statistical descriptors from the custom MobileNetV2 `xai_feature_layer`.",
        "Generate Med-MICN task probabilities, neural-symbolic probabilities, concept scores, and concept predictions.",
        "Fuse 26 statistical descriptors, 8 concept scores, and 2 Med-MICN diagnosis probabilities into 36 model features.",
        "Apply ZFMIS-style zero-fraction filtering and mutual-information ranking.",
        "Evaluate Decision Tree, RuleFit, stronger classifiers, and calibrated-threshold classifiers on saved workspace splits.",
        "Report rule artifacts, selected SFMOV/concept-aware SFMOV outputs, and implementation limitations.",
    ]
    for item in objectives:
        add_numbered(doc, item)


def chapter2(doc: Document) -> None:
    chapter(doc, "2", "Background")
    heading(doc, "2.1 Concept Overview")
    para(
        doc,
        "This chapter keeps the original background sequence but restricts it to the COVID-CT dataset source, "
        "the Med-MICN and SVIS-RULEX reference methods, and implementation-supported method/software sources. "
        "Unsupported broader literature claims are not used."
    )

    heading(doc, "2.1.1 Deep Learning and Feature Extraction for Medical Image Analysis", level=3)
    para(
        doc,
        "In the implemented SVIS-style branch, a custom MobileNetV2 model is used as the deep feature "
        "extractor. The base feature representation has 1280 dimensions after global average pooling, "
        "and the `xai_feature_layer` is a 128-dimensional dense representation. The statistical descriptors "
        "are computed over that 128-dimensional vector."
    )

    heading(doc, "2.1.2 Statistical Feature Engineering", level=3)
    para(
        doc,
        "The statistical branch follows the SVIS-RULEX idea of summarising deep activations using fixed "
        "statistical descriptors. In this implementation, exactly 26 descriptors are used."
    )
    add_table(
        doc,
        ["Descriptor group", "Implemented descriptors"],
        [
            ["Central tendency and extrema", "mean, median, min_value, max_value"],
            ["Dispersion", "std_dev, variance, range, iqr, mean_abs_dev, median_abs_dev, std_error_mean"],
            ["Percentiles", "percentile_25, percentile_50, percentile_75"],
            ["Shape and information", "skewness, kurtosis, entropy, shannon_entropy"],
            ["Energy and texture", "energy, contrast, root_mean_square, autocorrelation"],
            ["Ratios and means", "signal_to_noise, coef_of_var, harmonic_mean, geometric_mean"],
        ],
        "Table 2.1. Implemented 26 statistical descriptors.",
    )

    heading(doc, "2.1.3 Augmented ZFMIS (Zero-based Filtering with Mutual Importance Selection): Sparse-Value Filtering and Mutual-Information Ranking", level=3)
    para(
        doc,
        "ZFMIS in this implementation is applied as zero-fraction filtering followed by mutual-information "
        "ranking. A feature is retained if its zero fraction is less than or equal to 0.50 and removed if "
        "its zero fraction is greater than 0.50. The combined fused pool contains 36 candidate features; "
        "31 are retained after zero-fraction filtering."
    )

    heading(doc, "2.1.4 Rule-based Interpretability: Decision Trees and RuleFit", level=3)
    para(
        doc,
        "Decision Trees and RuleFit are used as interpretable rule-extraction models, consistent with the "
        "SVIS-RULEX background. The implementation exports Decision Tree text rules and RuleFit rule CSVs. "
        "For the implementation, RuleFit rules are exported from the RuleFit library and sorted by the "
        "library-provided `importance` column."
    )

    heading(doc, "2.1.5 Statistical Feature Map Overlay Visualisation (SFMOV) and Its Concept-Aware Extension", level=3)
    para(
        doc,
        "SFMOV is the visual explanation method used from the SVIS-RULEX background. The current workspace "
        "also produces concept-aware SFMOV overlays for selected test cases by modulating the statistical "
        "overlay with concept-score information. Concept Grad-CAM PNGs are not saved in the current workspace."
    )

    heading(doc, "2.1.6 Concept Bottleneck Models (CBMs) Concept-Based Interpretability", level=3)
    para(
        doc,
        "The Med-MICN branch provides concept-based evidence through concept labels, concept scores, concept "
        "embeddings, and neural-symbolic probabilities. The current SS-VIRULEX implementation should not be "
        "called a strict concept bottleneck model because the Med-MICN task head uses image features plus "
        "concept embeddings, and the final fused classifier also uses statistical descriptors and diagnosis "
        "probabilities."
    )

    heading(doc, "2.1.7 Large Multimodal Models (LMMs) and Vision-Language Models for Medical Concept Semantics", level=3)
    para(
        doc,
        "Med-MICN uses large multimodal and vision-language components in its reference-paper concept "
        "generation and annotation pipeline. In this workspace, runtime GPT-4V/BioViL generation is not "
        "confirmed. The implementation uses a completed COVID-CT concept CSV as materialized input."
    )

    heading(doc, "2.1.8 Neural-Symbolic Reasoning and Fuzzy Logic Rules", level=3)
    para(
        doc,
        "The Med-MICN reference method introduces neural-symbolic reasoning over learned concept information. "
        "The current workspace exports a neural-symbolic COVID probability as `med_neural_prob_covid`. This "
        "is separate from `med_task_prob_covid`, which is the task-head probability."
    )

    heading(doc, "2.1.9 Late Fusion of Heterogeneous Evidence Sources", level=3)
    para(
        doc,
        "Late fusion in the implemented SS-VIRULEX pipeline means row-aligned concatenation of statistical "
        "descriptors, concept scores, and Med-MICN diagnosis probabilities. It is not end-to-end deep joint "
        "training of all branches."
    )

    heading(doc, "2.1.10 Predictive Classifiers, Probability Calibration, and Validation-Based Threshold Selection", level=3)
    para(
        doc,
        "The stronger classifier family includes random forest, extra trees, gradient boosting, histogram "
        "gradient boosting, RBF SVM, and logistic regression. Calibration is applied only to these stronger "
        "classifiers. The calibrated stage uses sigmoid calibration with validation-selected thresholds."
    )

    heading(doc, "2.1.11 A Taxonomy of Explainability Modalities", level=3)
    para(
        doc,
        "The implemented explanation modalities are statistical feature descriptors, concept scores, Med-MICN "
        "task and neural-symbolic probabilities, ZFMIS rankings, Decision Tree rules, RuleFit rules, and "
        "selected SFMOV/concept-aware SFMOV overlays. Clinical validation and external validation are not "
        "implemented explanation modalities in this workspace."
    )

    heading(doc, "2.2 Literature Review")
    heading(doc, "2.2.1 Pillar I: Post-Hoc Visual and Attribution Methods", level=3)
    para(
        doc,
        "Within the allowed source boundary, this pillar is represented by the SVIS-RULEX SFMOV method. "
        "The current thesis does not claim experimental comparison against additional post-hoc XAI methods."
    )
    heading(doc, "2.2.2 Pillar II: Statistical and Rule-Based Interpretability Frameworks", level=3)
    para(
        doc,
        "SVIS-RULEX provides the reference background for statistical descriptors, ZFMIS, Decision Tree and "
        "RuleFit rule extraction, and SFMOV. The workspace adapts those components to the COVID-CT pipeline "
        "and to a fused statistical-concept feature pool."
    )
    heading(doc, "2.2.3 Pillar III: Concept-based and Neural-Symbolic Ante-Hoc Models", level=3)
    para(
        doc,
        "Med-MICN provides the reference background for concept prediction, concept embeddings, and "
        "neural-symbolic reasoning. The workspace uses the trained Med-MICN branch to export concept scores, "
        "thresholded concept predictions, task probability, and neural-symbolic probability."
    )
    heading(doc, "2.2.4 Pillar IV: Hybrid Evidence Fusion and Calibrated Clinical Decision Support", level=3)
    para(
        doc,
        "The implemented hybrid contribution is a calibrated classification and explanation pipeline, not a "
        "validated clinical decision-support system. It reports model metrics and explanation artifacts on "
        "the current COVID-CT split only."
    )


def chapter3(doc: Document) -> None:
    chapter(doc, "3", "Methodology")
    heading(doc, "3.1 Introduction to the Proposed Methodology")
    para(
        doc,
        "The methodology follows the existing thesis structure while correcting it to match the workspace. "
        "The pipeline prepares the COVID-CT split, performs train-only offline augmentation for the "
        "SVIS/statistical branch, extracts statistical descriptors, generates Med-MICN concept and diagnosis "
        "outputs, fuses the features, applies ZFMIS-style selection, evaluates rule and stronger classifiers, "
        "and produces selected visual explanations."
    )

    heading(doc, "3.2 SS-VIRULEX Framework Overview")
    para(
        doc,
        "SS-VIRULEX is implemented as a late-fusion framework. The statistical branch contributes 26 "
        "descriptors from the custom MobileNetV2 `xai_feature_layer`; the Med-MICN branch contributes 8 "
        "concept scores and 2 diagnosis probabilities. The resulting 36 model features are ranked and "
        "evaluated through rule-based and calibrated classifier stages."
    )
    add_figure(
        doc,
        ROOT / "ss_virulex_pipeline_diagram.png",
        "Figure 3.1. SS-VIRULEX implementation pipeline for COVID-CT classification and explanation.",
    )
    add_table(
        doc,
        ["Pipeline component", "Implemented workspace behavior"],
        [
            ["Dataset", "Official COVID-CT split: 746 images"],
            ["Augmentation", "Train-only offline augmentation for the SVIS/statistical branch"],
            ["Statistical branch", "Custom MobileNetV2 `xai_feature_layer`; 26 descriptors"],
            ["Concept branch", "Med-MICN concept scores, task probability, and neural-symbolic probability"],
            ["Fusion", "26 statistical + 8 concept + 2 diagnosis probability features"],
            ["Selection", "Zero-fraction filtering and mutual-information ranking"],
            ["Rules", "Decision Tree text rules and RuleFit rule CSVs"],
            ["Final comparison row", "Calibrated-threshold logistic regression on selected_9 fused features"],
        ],
        "Table 3.1. Corrected SS-VIRULEX implementation overview.",
    )

    heading(doc, "3.3 Dataset Preparation and Experimental Setting")
    para(
        doc,
        "The official COVID-CT split contains 746 images. The corrected split is 425 training images, "
        "118 validation images, and 203 test images. Validation and test images are not augmented."
    )
    add_table(
        doc,
        ["Split", "COVID", "NonCOVID", "Total"],
        [
            ["Train", "191", "234", "425"],
            ["Validation", "60", "58", "118"],
            ["Test", "98", "105", "203"],
            ["Total", "349", "397", "746"],
        ],
        "Table 3.2. Official COVID-CT split used in the workspace.",
    )
    para(
        doc,
        "For the SVIS/statistical branch, augmentation is applied only to the training split and saved as "
        "an offline augmented manifest. The augmented training manifest contains 936 rows, balanced as "
        "468 COVID and 468 NonCOVID rows. The augmentation uses mirror, small rotation within about +/-7 "
        "degrees, scale/crop/pad, brightness adjustment, and contrast adjustment. Vertical flips are not "
        "claimed because they are not confirmed by the audit or workspace."
    )

    heading(doc, "3.4 Shared Metadata Construction and Preprocessing")
    para(
        doc,
        "The SVIS/statistical branch uses 224 x 224 image processing for the custom MobileNetV2 pipeline. "
        "Med-MICN probability generation uses resize to 256, center crop to 224, and ImageNet-style "
        "normalization. These branch-specific preprocessing steps should not be collapsed into a single "
        "shared transform."
    )

    heading(doc, "3.5 Statistical Branch Based on SVIS-RULEX Principles")
    para(
        doc,
        "The statistical branch extracts the 128-dimensional `xai_feature_layer` representation from the "
        "custom MobileNetV2 model and computes the following 26 descriptors."
    )
    add_table(
        doc,
        ["No.", "Descriptor"],
        [[str(i + 1), name] for i, name in enumerate(DESCRIPTORS)],
        "Table 3.3. Implemented statistical descriptor columns.",
    )

    heading(doc, "3.6 Med-MICN Semantic and Neural-Symbolic Branch")
    para(
        doc,
        "The Med-MICN branch is not described as independently trained concept-only learning. The current "
        "training uses multiple losses involving task classification, concept prediction, neural-symbolic "
        "output, and concept-task prediction. The trained branch exports concept scores, thresholded concept "
        "predictions, task probability, and neural-symbolic probability."
    )
    add_table(
        doc,
        ["No.", "Implemented concept"],
        [[str(i + 1), name] for i, name in enumerate(CONCEPTS)],
        "Table 3.4. Actual Med-MICN concept vocabulary used in the workspace.",
    )
    para(
        doc,
        "`med_task_prob_covid` is the probability from the Med-MICN task classification head. "
        "`med_neural_prob_covid` is the probability from the Med-MICN neural-symbolic reasoning output. "
        "These probabilities are distinct and are both included in the fused feature set."
    )

    heading(doc, "3.7 Statistical-Semantic Feature Fusion")
    para(
        doc,
        "The fused model feature vector is defined as: fused_features = 26 statistical descriptors + 8 "
        "concept scores + `med_task_prob_covid` + `med_neural_prob_covid`. Therefore the implemented model "
        "feature dimension is 36. The full `fused_features_covid.csv` file contains extra identifier, split, "
        "label, and path metadata columns in addition to the 36 model features."
    )

    heading(doc, "3.8 Augmented ZFMIS Feature Selection")
    para(
        doc,
        "The combined fused pool contains 36 candidate model features. Zero-fraction filtering retains a "
        "feature if `zero_fraction <= 0.50` and removes it if `zero_fraction > 0.50`. After this filtering, "
        "31 features are retained. Mutual-information ranking is then applied to the retained features, and "
        "selected feature sets are saved for selected_18, selected_15, selected_12, selected_9, selected_6, "
        "and selected_3."
    )

    heading(doc, "3.9 Rule-Based Classification and Explanation")
    para(
        doc,
        "The combined Decision Tree is not a fixed entropy/depth-10 model. It is selected by grid search "
        "over criterion `gini` and `entropy`; max depth 2, 3, 4, 5, 8, and None; min samples leaf 1, 2, 4, "
        "and 8; and min samples split 2, 5, and 10. The best combined Decision Tree uses the selected_15 "
        "feature set, Gini criterion, max depth 3, min samples leaf 4, and min samples split 2."
    )
    para(
        doc,
        "Combined RuleFit is run only on selected_3 and selected_6. Its grid uses tree size 3 or 4, sample "
        "fraction 0.8, and max rules 50 or 100. The selected_3 configuration is best by test balanced "
        "accuracy, while selected_6 has the higher validation balanced accuracy. This distinction is reported "
        "as a methodological caveat."
    )

    heading(doc, "3.10 Calibrated Stronger Classifiers and Validation-Based Threshold Selection")
    para(
        doc,
        "The implemented stronger classifier family consists of random forest, extra trees, gradient boosting, "
        "histogram gradient boosting, RBF SVM, and logistic regression. Calibration applies only to this "
        "stronger classifier family. The calibrated stage uses sigmoid calibration and selects the operating "
        "threshold on the validation split using balanced accuracy, with weighted F1 and accuracy as tie-breakers."
    )

    heading(doc, "3.11 Visual Explanation Workflow")
    para(
        doc,
        "The visual workflow uses SFMOV and concept-aware SFMOV only. Concept-aware SFMOV uses selected test "
        "cases and a scalar concept modulation based on concept-score information. Concept Grad-CAM PNGs are "
        "not saved in the current workspace, so this thesis does not claim concept Grad-CAM as an implemented "
        "output."
    )

    heading(doc, "3.12 Implementation Details, and Reproducibility")
    para(
        doc,
        "The main pipeline evidence is the workspace implementation and outputs under `03_SS-VIRULEX-outputs`. "
        "Key artifacts include `fused_features_covid.csv`, `covid_statistical_features.csv`, "
        "`covid_concept_probabilities.csv`, `zfmis_ranking_covid.csv`, Decision Tree and RuleFit rule files, "
        "`metrics_covid.csv`, `metrics_three_model_comparison_covid.csv`, and `heatmap_cases_covid.csv`."
    )
    para(
        doc,
        "The thesis reports current implementation limitations: runtime GPT-4V/BioViL generation scripts are "
        "not confirmed; concept Grad-CAM PNGs are not saved; RuleFit per-image predictions and a final "
        "serialized RuleFit model object are not found in consolidated outputs; no clinical validation, "
        "reader study, or external validation dataset is present."
    )


def chapter4(doc: Document) -> None:
    chapter(doc, "4", "Testing and Results")
    heading(doc, "4.1 Dataset and Experimental Setup")
    heading(doc, "4.1.1 Dataset Identity and Official Split", level=3)
    para(
        doc,
        "The corrected dataset setup uses the official 746-image COVID-CT split. The training split contains "
        "425 images, the validation split contains 118 images, and the test split contains 203 images."
    )
    add_table(
        doc,
        ["Split", "COVID", "NonCOVID", "Total"],
        [["Train", "191", "234", "425"], ["Validation", "60", "58", "118"], ["Test", "98", "105", "203"]],
        "Table 4.1. Official split used for testing and results.",
    )
    heading(doc, "4.1.2 Augmented Processing Manifest", level=3)
    para(
        doc,
        "The augmented processing manifest has 1257 rows overall: 936 augmented training rows, 118 validation "
        "rows, and 203 test rows. The augmented training rows are balanced as 468 COVID and 468 NonCOVID."
    )

    heading(doc, "4.2 Branch-Level Outputs")
    heading(doc, "4.2.1 Statistical Branch Output and Custom MobileNetV2 Metrics", level=3)
    para(
        doc,
        "The custom MobileNetV2 branch is a feature-extraction and neural baseline branch. It is not the "
        "final interpretable SS-VIRULEX operating model."
    )
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ["Accuracy", "0.7537"],
            ["Balanced accuracy", "0.7520"],
            ["Weighted precision", "0.7546"],
            ["Weighted recall", "0.7537"],
            ["Weighted F1", "0.7530"],
            ["ROC AUC", "0.7997"],
            ["Confusion matrix", "[[84, 21], [29, 69]]"],
            ["Best hyperparameters", "learning_rate 0.001; dense_units 256; dropout 0.5"],
            ["Epochs run", "7"],
        ],
        "Table 4.2. Custom MobileNetV2 test metrics.",
    )
    heading(doc, "4.2.2 Med-MICN Concept and Neural-Symbolic Branch", level=3)
    add_table(
        doc,
        ["Output", "Accuracy", "F1", "Recall / balanced accuracy", "ROC AUC"],
        [
            ["Task head", "0.7044", "0.7043", "0.7044", "0.7598"],
            ["Neural-symbolic", "0.7143", "0.7106", "0.7112", "0.7459"],
            ["Concept prediction", "0.7217", "0.6591 macro", "-", "-"],
        ],
        "Table 4.3. Med-MICN branch test metrics.",
    )
    para(
        doc,
        "The neural-symbolic output has higher macro recall/balanced accuracy than the task head, while the "
        "task head has the higher ROC AUC. These values should not be interpreted as a clinical validation."
    )
    add_table(
        doc,
        ["Concept", "Accuracy", "F1", "Recall"],
        [
            ["absence_of_lobar_consolidation", "0.9655", "0.9825", "0.9949"],
            ["increased_density_in_the_lung", "0.6946", "0.8198", "1.0000"],
            ["localized_or_diffuse_presentation", "0.6897", "0.8163", "1.0000"],
            ["multilobar_distribution", "0.6798", "0.8094", "1.0000"],
            ["peripheral_ground_glass_opacities", "0.5764", "0.7313", "0.9915"],
            ["bilateral_involvement", "0.6108", "0.7285", "0.8983"],
            ["ground_glass_appearance", "0.6355", "0.7218", "0.7059"],
            ["crazy_paving_pattern", "0.9212", "0.0000", "0.0000"],
        ],
        "Table 4.4. Per-concept Med-MICN test metrics.",
    )

    heading(doc, "4.3 Fused Feature Construction and ZFMIS Results")
    heading(doc, "4.3.1 Late Fusion of Heterogeneous Evidence", level=3)
    para(
        doc,
        "The fused table contains 1257 rows and 42 total columns. The model feature pool has 36 features: "
        "26 statistical descriptors, 8 concept scores, and 2 Med-MICN diagnosis probabilities. The remaining "
        "columns are metadata."
    )
    heading(doc, "4.3.2 Augmented ZFMIS Feature Ranking and Selection", level=3)
    para(
        doc,
        "The ZFMIS ranking export contains all 36 candidate features and marks selected features. The "
        "selected feature sets are drawn from the 31 retained features after zero-fraction filtering."
    )
    add_table(
        doc,
        ["Rank", "Feature", "Type", "Mutual information"],
        [
            ["1", "med_task_prob_covid", "Med-MICN diagnosis probability", "0.4739"],
            ["2", "med_neural_prob_covid", "Med-MICN diagnosis probability", "0.4678"],
            ["3", "concept_score_absence_of_lobar_consolidation", "Concept score", "0.4595"],
            ["4", "concept_score_ground_glass_appearance", "Concept score", "0.4584"],
            ["5", "concept_score_crazy_paving_pattern", "Concept score", "0.4410"],
            ["6", "concept_score_bilateral_involvement", "Concept score", "0.2946"],
            ["7", "concept_score_localized_or_diffuse_presentation", "Concept score", "0.0664"],
            ["8", "entropy", "Statistical descriptor", "0.0437"],
            ["9", "shannon_entropy", "Statistical descriptor", "0.0437"],
            ["10", "concept_score_multilobar_distribution", "Concept score", "0.0408"],
        ],
        "Table 4.5. Top ZFMIS-ranked fused features.",
    )

    heading(doc, "4.4 Rule-Based Explanation Results")
    heading(doc, "4.4.1 Best Combined Decision Tree", level=3)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Feature set", "selected_15"],
            ["Best parameters", "criterion gini; max_depth 3; min_samples_leaf 4; min_samples_split 2"],
            ["CV balanced accuracy", "0.9113"],
            ["Test accuracy", "0.6700"],
            ["Test balanced accuracy", "0.6741"],
            ["Weighted F1", "0.6656"],
            ["ROC AUC", "0.7349"],
            ["Confusion matrix", "[[58, 47], [20, 78]]"],
        ],
        "Table 4.6. Best combined Decision Tree result.",
    )
    heading(doc, "4.4.2 Best Combined RuleFit", level=3)
    add_table(
        doc,
        ["Feature set", "Best params", "Validation balanced accuracy", "Test balanced accuracy", "Weighted F1", "Nonzero rules"],
        [
            ["selected_3", "tree_size 3; sample_fract 0.8; max_rules 50", "0.7273", "0.6830", "0.6775", "16"],
            ["selected_6", "tree_size 3; sample_fract 0.8; max_rules 100", "0.7615", "0.6782", "0.6723", "32"],
        ],
        "Table 4.7. Combined RuleFit results and selection caveat.",
    )
    para(
        doc,
        "The selected_3 RuleFit result is best by test balanced accuracy. The selected_6 result has higher "
        "validation balanced accuracy. This caveat is reported because validation-based selection and "
        "test-based ranking lead to different preferred rows."
    )

    heading(doc, "4.5 Visual Explanation Results")
    para(
        doc,
        "The visual outputs are SFMOV and concept-aware SFMOV only. The workspace saves selected heatmap "
        "cases and does not save concept Grad-CAM PNGs. The visual examples are qualitative explanation "
        "artifacts and are not clinical validation."
    )
    add_figure(
        doc,
        ROOT / "03_SS-VIRULEX-outputs/visual_explanations/heatmaps/ss_virulex_concept_aware/covid_bmj_m606_full_p4_22_2_concept_aware.png",
        "Figure 4.1. Example concept-aware SFMOV overlay for a COVID test image.",
    )
    add_figure(
        doc,
        ROOT / "03_SS-VIRULEX-outputs/visual_explanations/heatmaps/ss_virulex_concept_aware/noncovid_849_concept_aware.png",
        "Figure 4.2. Example concept-aware SFMOV overlay for a NonCOVID test image.",
    )

    heading(doc, "4.6 Final Output Reporting")
    para(
        doc,
        "Final reporting is based on saved workspace artifacts. The predictions file contains Decision Tree "
        "and final calibrated-threshold predictions; RuleFit per-image prediction columns are blank in the "
        "consolidated prediction file. The RuleFit rules are saved as rule CSVs, but a final serialized "
        "RuleFit model object is not found in the consolidated outputs."
    )
    add_table(
        doc,
        ["Artifact", "Purpose"],
        [
            ["fused_features_covid.csv", "Fused statistical, concept, and diagnosis-probability features"],
            ["covid_statistical_features.csv", "26 statistical descriptors for the augmented manifest"],
            ["covid_concept_probabilities.csv", "Med-MICN concept scores and diagnosis probabilities"],
            ["zfmis_ranking_covid.csv", "Combined fused feature ranking"],
            ["06_combined_best_decision_tree_rules.txt", "Best Decision Tree text rules"],
            ["06_combined_best_rulefit_rules.csv", "Best RuleFit rule table"],
            ["metrics_three_model_comparison_covid.csv", "Final three-row model comparison"],
            ["predictions_covid.csv", "Consolidated test predictions"],
            ["heatmap_cases_covid.csv", "Selected visual explanation cases"],
        ],
        "Table 4.8. Major saved output artifacts.",
    )

    heading(doc, "4.7 Final Comparative Results")
    para(
        doc,
        "The final SS-VIRULEX comparison row is calibrated-threshold logistic regression using the selected_9 "
        "fused feature set. Decision Tree and RuleFit are rule-extraction/explanation artifacts and comparison "
        "models; they are not the final best predictive row."
    )
    add_table(
        doc,
        ["Rank", "Model family", "Selected model", "Feature set", "Accuracy", "Balanced accuracy", "Weighted F1", "ROC AUC"],
        [
            ["1", "SS-VIRULEX", "calibrated-threshold logistic regression", "selected_9", "0.7340", "0.7313", "0.7321", "0.7591"],
            ["2", "Med-MICN", "neural-symbolic", "neural_symbolic", "0.7143", "0.7112", "0.7106", "0.7459"],
            ["3", "SVIS-RULEX", "Decision Tree", "selected_6", "0.5616", "0.5544", "0.5415", "0.5475"],
        ],
        "Table 4.9. Final three-model comparison.",
    )
    para(
        doc,
        "This table is the final three-row comparison saved by the workspace. It is not a comparison against "
        "all external XAI methods and should not be described as clinical validation."
    )


def chapter5(doc: Document) -> None:
    chapter(doc, "5", "Conclusion")
    para(
        doc,
        "This thesis implemented and revised SS-VIRULEX as a grounded COVID-CT explainability pipeline. "
        "The implemented feature fusion combines 26 statistical descriptors, 8 Med-MICN concept scores, "
        "and 2 Med-MICN diagnosis probabilities into 36 model features. ZFMIS-style filtering retains "
        "31 of the 36 candidate features before mutual-information ranking and selected feature-set evaluation."
    )
    para(
        doc,
        "The final SS-VIRULEX comparison row is calibrated-threshold logistic regression using selected_9 "
        "fused features, with test accuracy 0.7340, balanced accuracy 0.7313, weighted F1 0.7321, and ROC "
        "AUC 0.7591. Decision Tree and RuleFit provide rule-extraction artifacts and comparison results, "
        "but they are not the final best predictive row."
    )
    para(
        doc,
        "The implemented explanation artifacts include ZFMIS feature rankings, Decision Tree rules, RuleFit "
        "rules, and selected SFMOV/concept-aware SFMOV visual overlays. The current workspace does not "
        "support claims of clinical validation, clinician trust evaluation, external validation, or deployment "
        "readiness."
    )


def chapter6(doc: Document) -> None:
    chapter(doc, "6", "Future Work")
    para(
        doc,
        "Future work should address the implementation gaps identified in the audit without overstating the "
        "current results. First, runtime concept-label generation and annotation should be made reproducible "
        "if the Med-MICN large multimodal and vision-language stages are to be claimed as active runtime "
        "components. Second, concept-specific saliency or Grad-CAM-style outputs should be generated and "
        "saved if future versions claim full concept-spatial alignment."
    )
    para(
        doc,
        "The RuleFit stage should save a final model object and per-image RuleFit predictions. Additional "
        "experiments should document systematic ablations over statistical-only, concept-only, diagnosis-only, "
        "and fused feature groups using clear validation-based selection rules. External validation and "
        "clinical reader studies would be required before making clinical relevance, trust, or deployment "
        "claims."
    )


def appendix_and_refs(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Appendix", level=1)
    heading(doc, "A Lists", level=2)
    heading(doc, "List of Abbreviations", level=3)
    add_table(
        doc,
        ["Abbreviation", "Meaning"],
        [
            ["SS-VIRULEX", "Semantic-Statistical Visual Interpretable Rule-based Explainer"],
            ["SVIS-RULEX", "Statistical, visual, and rule-based explainable AI framework"],
            ["Med-MICN", "Multi-dimensional explanation alignment method for medical classification"],
            ["ZFMIS", "Zero-based filtering with mutual importance selection"],
            ["SFMOV", "Statistical Feature Map Overlay Visualisation"],
            ["DT", "Decision Tree"],
            ["ROC AUC", "Area under the receiver operating characteristic curve"],
        ],
        "Table A.1. Abbreviations used in the revised thesis.",
    )
    heading(doc, "List of Figures", level=3)
    doc.add_paragraph("Figure 3.1. SS-VIRULEX implementation pipeline for COVID-CT classification and explanation.")
    doc.add_paragraph("Figure 4.1. Example concept-aware SFMOV overlay for a COVID test image.")
    doc.add_paragraph("Figure 4.2. Example concept-aware SFMOV overlay for a NonCOVID test image.")
    heading(doc, "List of Tables", level=3)
    for table_name in [
        "Table 2.1. Implemented 26 statistical descriptors.",
        "Table 3.1. Corrected SS-VIRULEX implementation overview.",
        "Table 3.2. Official COVID-CT split used in the workspace.",
        "Table 3.3. Implemented statistical descriptor columns.",
        "Table 3.4. Actual Med-MICN concept vocabulary used in the workspace.",
        "Table 4.1. Official split used for testing and results.",
        "Table 4.2. Custom MobileNetV2 test metrics.",
        "Table 4.3. Med-MICN branch test metrics.",
        "Table 4.4. Per-concept Med-MICN test metrics.",
        "Table 4.5. Top ZFMIS-ranked fused features.",
        "Table 4.6. Best combined Decision Tree result.",
        "Table 4.7. Combined RuleFit results and selection caveat.",
        "Table 4.8. Major saved output artifacts.",
        "Table 4.9. Final three-model comparison.",
        "Table A.1. Abbreviations used in the revised thesis.",
    ]:
        doc.add_paragraph(table_name)

    doc.add_page_break()
    doc.add_heading("References", level=1)
    references = [
        (
            "Yang, X., He, X., Zhao, J., Zhang, Y., Zhang, S., & Xie, P. (2020). "
            "COVID-CT-Dataset: A CT scan dataset about COVID-19. arXiv. "
            "https://doi.org/10.48550/arXiv.2003.13865"
        ),
        (
            "Hu, L., Lai, S., Chen, W., Xiao, H., Lin, H., Yu, L., Zhang, J., & Wang, D. "
            "(2024). Towards multi-dimensional explanation alignment for medical classification. "
            "In Advances in Neural Information Processing Systems 37. "
            "https://doi.org/10.52202/079017-4119"
        ),
        (
            "Ullah, N., Guzman-Aroca, F., Martinez-Alvarez, F., De Falco, I., & Sannino, G. "
            "(2025). A novel explainable AI framework for medical image classification integrating "
            "statistical, visual, and rule-based methods. Medical Image Analysis, 105, Article 103665. "
            "https://doi.org/10.1016/j.media.2025.103665"
        ),
        (
            "Sandler, M., Howard, A., Zhu, M., Zhmoginov, A., & Chen, L.-C. (2018). "
            "MobileNetV2: Inverted residuals and linear bottlenecks. In Proceedings of the IEEE/CVF "
            "Conference on Computer Vision and Pattern Recognition (pp. 4510-4520). "
            "https://doi.org/10.1109/CVPR.2018.00474"
        ),
        (
            "He, K., Zhang, X., Ren, S., & Sun, J. (2016). Deep residual learning for image recognition. "
            "In Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition "
            "(pp. 770-778). https://doi.org/10.1109/CVPR.2016.90"
        ),
        (
            "Friedman, J. H., & Popescu, B. E. (2008). Predictive learning via rule ensembles. "
            "The Annals of Applied Statistics, 2(3), 916-954. https://doi.org/10.1214/07-AOAS148"
        ),
        (
            "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., "
            "Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., "
            "Cournapeau, D., Brucher, M., Perrot, M., & Duchesnay, E. (2011). Scikit-learn: "
            "Machine learning in Python. Journal of Machine Learning Research, 12, 2825-2830."
        ),
        (
            "Paszke, A., Gross, S., Massa, F., Lerer, A., Bradbury, J., Chanan, G., Killeen, T., "
            "Lin, Z., Gimelshein, N., Antiga, L., Desmaison, A., Kopf, A., Yang, E., DeVito, Z., "
            "Raison, M., Tejani, A., Chilamkurthy, S., Steiner, B., Fang, L., ... Chintala, S. "
            "(2019). PyTorch: An imperative style, high-performance deep learning library. "
            "In Advances in Neural Information Processing Systems 32."
        ),
        (
            "Abadi, M., Agarwal, A., Barham, P., Brevdo, E., Chen, Z., Citro, C., Corrado, G. S., "
            "Davis, A., Dean, J., Devin, M., Ghemawat, S., Goodfellow, I., Harp, A., Irving, G., "
            "Isard, M., Jia, Y., Jozefowicz, R., Kaiser, L., Kudlur, M., ... Zheng, X. (2016). "
            "TensorFlow: Large-scale machine learning on heterogeneous systems. https://www.tensorflow.org/"
        ),
    ]
    for reference in references:
        doc.add_paragraph(reference)


def build_docx() -> None:
    doc = Document()
    configure_doc(doc)
    cover(doc)
    contents(doc)
    chapter1(doc)
    chapter2(doc)
    chapter3(doc)
    chapter4(doc)
    chapter5(doc)
    chapter6(doc)
    appendix_and_refs(doc)
    doc.save(OUT_DOCX)


def build_changelog() -> None:
    text = """# SS-VIRULEX Revised Thesis Change Log

## Major Corrections Applied

- Corrected the official COVID-CT split to 746 images: 425 train, 118 validation, and 203 test.
- Corrected augmentation as train-only offline augmentation for the SVIS/statistical branch, with 936 augmented training rows balanced as 468 COVID and 468 NonCOVID.
- Replaced the statistical descriptor list with the exact 26 implemented descriptors from the workspace.
- Replaced the concept vocabulary with the exact eight Med-MICN workspace concepts.
- Rewrote the Med-MICN branch as multi-loss training involving task, concept, neural-symbolic, and concept-task losses.
- Distinguished `med_task_prob_covid` from `med_neural_prob_covid`.
- Corrected feature fusion to 26 statistical descriptors, 8 concept scores, and 2 Med-MICN diagnosis probabilities for 36 model features.
- Corrected ZFMIS to retain `zero_fraction <= 0.50`, remove `zero_fraction > 0.50`, and report 36 candidate features with 31 retained after filtering.
- Corrected Decision Tree configuration to the implemented grid search and reported the selected_15 best combined Decision Tree parameters.
- Corrected RuleFit configuration to selected_3 and selected_6 only, with the required validation/test selection caveat.
- Corrected the stronger classifier family to random forest, extra trees, gradient boosting, histogram gradient boosting, RBF SVM, and logistic regression.
- Corrected calibration as sigmoid calibration with validation-selected thresholds for the stronger classifier family only.
- Corrected the final SS-VIRULEX comparison row to calibrated-threshold logistic regression using selected_9 fused features.
- Restricted visual explanation claims to SFMOV and concept-aware SFMOV.
- Added a professional SS-VIRULEX pipeline figure for Chapter 3.

## Sections Revised

- Chapter 1: Motivation, Problem Statement, and Objectives were softened and grounded in the implemented workspace.
- Chapter 2: Background was restricted to Med-MICN 2024, SVIS-RULEX 2025, and implementation-supported terminology.
- Chapter 3: Methodology was corrected for dataset split, augmentation, descriptors, concepts, fusion, ZFMIS, classifiers, calibration, and visual outputs.
- Chapter 4: Results were revised to match saved workspace metrics and artifacts.
- Chapter 5: Conclusion was filled with implementation-supported findings only.
- Chapter 6: Future Work was filled with audit-supported implementation gaps.
- Appendix and References were updated with abbreviations, tables, figures, the COVID-CT dataset source, the two core reference methods, and implementation-supported method/software sources.

## Unsupported Claims Removed or Softened

- Removed unsupported claims of clinical validation, clinician trust study, clinical deployment readiness, external validation datasets, and comparison against many other state-of-the-art XAI methods.
- Qualified GPT-4V/BioViL as Med-MICN reference-paper background/upstream context, not a confirmed runtime stage in the current workspace.
- Removed unsupported implemented-descriptor claims for spectral entropy, spectral energy, and zero crossing rate.
- Removed XGBoost from the implemented classifier list.
- Removed or qualified concept Grad-CAM claims because concept Grad-CAM PNGs are not saved in the current workspace.

## Limitations Added

- Runtime GPT-4V/BioViL generation is not confirmed in the current workspace.
- Concept Grad-CAM PNGs are not saved.
- RuleFit per-image predictions and a final serialized RuleFit model object are not found in consolidated outputs.
- The final best SS-VIRULEX predictive row is calibrated logistic regression, not a rule model.
- No clinical validation, reader study, external validation, or deployment-readiness evidence is present.

## Audit Items Not Applied

- Section 9 of `THESIS_IMPLEMENTATION_REVISION_AUDIT.md` was intentionally skipped. The thesis was not restructured according to the recommended corrected thesis structure; the existing draft structure was preserved as much as possible.
- Full visual render QA with LibreOffice could not be completed because `soffice` is not installed in the environment. Structural DOCX creation and text verification were completed.
"""
    OUT_CHANGELOG.write_text(text, encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_changelog()
    print(OUT_DOCX)
    print(OUT_CHANGELOG)
